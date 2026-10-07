"""Read the contractor's delivery-plan document (.docx or .doc) into a neutral `RawDoc`.

No database and no HTTP here. `.docx` goes through python-docx; the old binary `.doc` goes through
the `antiword` program (DocBook output keeps every paragraph of a table cell). Anything that is not
a Word file, is too big, or looks like a decompression bomb is refused with a readable message.
"""

import io
import os
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field

import docx
from defusedxml import ElementTree
from docx.table import Table
from docx.text.paragraph import Paragraph

MAX_BYTES = 5 * 1024 * 1024
MAX_UNPACKED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
ANTIWORD_TIMEOUT_SECONDS = 30

_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class PlanFileError(ValueError):
    """The file cannot be read as a plan document; the message is shown to the user."""


@dataclass
class RawTable:
    """Rows of cells; a cell is its paragraphs joined by newlines. A row spanning the whole table
    (a stage heading) has a single cell."""

    rows: list[list[str]] = field(default_factory=list)


@dataclass
class RawDoc:
    """Paragraph texts and tables in document order."""

    blocks: list[str | RawTable] = field(default_factory=list)

    @property
    def tables(self) -> list[RawTable]:
        return [b for b in self.blocks if isinstance(b, RawTable)]

    @property
    def lines(self) -> list[str]:
        return [b for b in self.blocks if isinstance(b, str)]


# --- .docx ---------------------------------------------------------------------------------------


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNPACKED_BYTES:
                raise PlanFileError("Tệp Word có cấu trúc bất thường nên không được đọc")
            names = {i.filename for i in infos}
    except zipfile.BadZipFile as exc:
        raise PlanFileError("Không đọc được tệp Word (.docx)") from exc
    if "word/document.xml" not in names:
        raise PlanFileError("Tệp không phải văn bản Word (.docx)")
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise PlanFileError("Tệp Word có macro, hãy lưu lại dưới dạng .docx không macro")


def _from_docx(data: bytes) -> RawDoc:
    _check_zip(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises many types for damaged files
        raise PlanFileError("Không đọc được tệp Word (.docx)") from exc
    raw = RawDoc()
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document).text.strip()
            if text:
                raw.blocks.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            rows: list[list[str]] = []
            for row in table.rows:
                cells: list[str] = []
                seen: set[int] = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:  # a merged cell is reported once per column it spans
                        continue
                    seen.add(id(cell._tc))
                    cells.append(
                        "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                    )
                rows.append(cells)
            raw.blocks.append(RawTable(rows))
    return raw


# --- .doc (antiword) -----------------------------------------------------------------------------


def _antiword() -> str:
    found = os.environ.get("ANTIWORD_BIN") or shutil.which("antiword")
    if not found:
        raise PlanFileError(
            "Máy chủ chưa đọc được tệp .doc cũ, hãy lưu tệp dưới dạng .docx rồi tải lên lại"
        )
    return found


def _text(el: object) -> str:
    return "".join(el.itertext())  # type: ignore[attr-defined]


def _entry_text(entry: object) -> str:
    lines = [ln.strip() for ln in _text(entry).splitlines()]
    return "\n".join(ln for ln in lines if ln)


def _from_doc(data: bytes) -> RawDoc:
    binary = _antiword()
    with tempfile.NamedTemporaryFile(suffix=".doc") as tmp:
        tmp.write(data)
        tmp.flush()
        try:
            done = subprocess.run(  # noqa: S603 - fixed argv, the file is a temporary path
                [binary, "-m", "UTF-8.txt", "-x", "db", tmp.name],
                capture_output=True,
                timeout=ANTIWORD_TIMEOUT_SECONDS,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise PlanFileError("Đọc tệp .doc quá lâu, hãy lưu dưới dạng .docx") from exc
        except OSError as exc:  # the program is missing or cannot run
            raise PlanFileError(
                "Máy chủ chưa đọc được tệp .doc cũ, hãy lưu tệp dưới dạng .docx rồi tải lên lại"
            ) from exc
    if done.returncode != 0 or not done.stdout.strip():
        raise PlanFileError("Không đọc được tệp Word (.doc), hãy lưu dưới dạng .docx")
    try:
        root = ElementTree.fromstring(done.stdout)
    except ElementTree.ParseError as exc:
        raise PlanFileError("Không đọc được tệp Word (.doc), hãy lưu dưới dạng .docx") from exc

    raw = RawDoc()
    parents = {child: parent for parent in root.iter() for child in parent}

    def inside_table(el: object) -> bool:
        cur = parents.get(el)
        while cur is not None:
            if cur.tag in {"entry", "informaltable", "table"}:
                return True
            cur = parents.get(cur)
        return False

    for el in root.iter():
        if el.tag in {"informaltable", "table"} and not inside_table(el):
            rows = [[_entry_text(entry) for entry in row.iter("entry")] for row in el.iter("row")]
            raw.blocks.append(RawTable(rows))
        elif el.tag == "para" and not inside_table(el) and el.find(".//informaltable") is None:
            for line in _text(el).splitlines():
                if line.strip():
                    raw.blocks.append(" ".join(line.split()))
    return raw


# --- entry point ---------------------------------------------------------------------------------


def read_plan_document(data: bytes, filename: str) -> RawDoc:
    if len(data) > MAX_BYTES:
        raise PlanFileError("Tệp quá lớn (tối đa 5 MB)")
    name = filename.lower()
    if not name.endswith((".docx", ".doc")):
        raise PlanFileError("Chỉ nhận tệp Word (.docx hoặc .doc)")
    if data.startswith(_ZIP_MAGIC):
        raw = _from_docx(data)
    elif data.startswith(_OLE_MAGIC):
        raw = _from_doc(data)
    else:
        raise PlanFileError("Tệp không phải văn bản Word")
    if not raw.tables:
        raise PlanFileError("Không thấy bảng nào trong tệp kế hoạch")
    return raw
