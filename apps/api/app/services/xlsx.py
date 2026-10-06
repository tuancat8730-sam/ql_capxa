"""Excel output shared by the list exports (SPEC 4.13): Vietnamese headers, real dates and money."""

import io
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Literal
from urllib.parse import quote

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HOT_FILL = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
HOT_FONT = Font(color="9C0006", bold=True)
HEAD_FILL = PatternFill("solid", start_color="D9E2F3", end_color="D9E2F3")
TOTAL_FILL = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

Kind = Literal["text", "date", "datetime", "money", "int", "pct"]
_FORMATS: dict[str, str | None] = {
    "text": None,
    "date": "DD/MM/YYYY",
    "datetime": "DD/MM/YYYY HH:MM",
    "money": "#,##0",
    "int": "0",
    "pct": "0.0",
}


@dataclass(frozen=True)
class Col:
    header: str
    kind: Kind = "text"
    width: int = 16


@dataclass
class Row:
    values: Sequence[Any]
    hot: frozenset[int] = field(default_factory=frozenset)  # column indexes shown in red


def _plain(value: Any, kind: Kind) -> Any:
    """Excel wants naive datetimes and plain numbers."""
    if value is None:
        return None
    if kind == "datetime" and isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if kind == "date" and isinstance(value, datetime):
        return value.date()
    if kind in {"money", "int", "pct"}:
        return float(value) if kind == "pct" else int(value)
    return value


def build_sheet(
    sheet: str,
    title: str,
    subtitle: str,
    columns: Sequence[Col],
    rows: Sequence[Row],
    totals: dict[int, Any] | None = None,
) -> bytes:
    """One worksheet: title, subtitle, header row, data rows and an optional totals row.

    `totals` maps a column index to its total (the first column then reads "Tổng cộng").
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet[:31]
    last = get_column_letter(len(columns))
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(f"A1:{last}1")
    ws["A2"] = subtitle
    ws.merge_cells(f"A2:{last}2")

    head = 4
    for c, col in enumerate(columns, start=1):
        cell = ws.cell(row=head, column=c, value=col.header)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = col.width
    ws.row_dimensions[head].height = 34

    for r, row in enumerate(rows, start=head + 1):
        for c, (col, value) in enumerate(zip(columns, row.values, strict=True), start=1):
            cell = ws.cell(row=r, column=c, value=_plain(value, col.kind))
            fmt = _FORMATS[col.kind]
            if fmt:
                cell.number_format = fmt
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col.kind == "text")
            if c - 1 in row.hot:
                cell.fill, cell.font = HOT_FILL, HOT_FONT

    if totals is not None:
        r = head + 1 + len(rows)
        for c, col in enumerate(columns, start=1):
            value = "Tổng cộng" if c == 1 else totals.get(c - 1)
            cell = ws.cell(row=r, column=c, value=_plain(value, col.kind) if c > 1 else value)
            cell.font = Font(bold=True)
            cell.fill = TOTAL_FILL
            cell.border = BORDER
            fmt = _FORMATS[col.kind]
            if fmt and c > 1:
                cell.number_format = fmt

    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(content: bytes, filename: str) -> Response:
    disposition = f"attachment; filename=\"{filename}\"; filename*=UTF-8''{quote(filename)}"
    return Response(content=content, media_type=XLSX, headers={"Content-Disposition": disposition})


def as_date(value: date | datetime | None) -> date | None:
    return value.date() if isinstance(value, datetime) else value
