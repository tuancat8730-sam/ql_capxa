"""Text extraction for search (SPEC 4.6, section 9). No OCR: scanned PDFs are flagged instead."""

import io
from dataclasses import dataclass
from typing import Literal

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

Status = Literal["done", "needs_ocr", "unsupported", "failed"]

MAX_TEXT_CHARS = 2_000_000
# A PDF whose pages yield almost no characters is an image scan, not a text document.
MIN_PDF_CHARS = 20
MAX_XLSX_CELLS = 200_000

_TEXT_TYPES = {"text/plain", "text/csv", "text/markdown"}
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class ExtractionResult:
    status: Status
    text: str | None = None


def _clean(text: str) -> str:
    return " ".join(text.replace("\x00", " ").split())[:MAX_TEXT_CHARS]


def _pdf(data: bytes) -> ExtractionResult:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        return ExtractionResult("failed")
    pages = (page.extract_text() or "" for page in reader.pages)
    text = _clean(" ".join(pages))
    if len(text) < MIN_PDF_CHARS:
        return ExtractionResult("needs_ocr")
    return ExtractionResult("done", text)


def _docx(data: bytes) -> ExtractionResult:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return ExtractionResult("done", _clean(" ".join(parts)))


def _xlsx(data: bytes) -> ExtractionResult:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    seen = 0
    for sheet in workbook.worksheets:
        parts.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    parts.append(str(value))
                    seen += 1
            if seen >= MAX_XLSX_CELLS:
                break
    return ExtractionResult("done", _clean(" ".join(parts)))


def extract_text(data: bytes, mime_type: str, file_name: str = "") -> ExtractionResult:
    """Never raises: a damaged file becomes `failed` so the upload itself is not lost."""
    name = file_name.lower()
    try:
        if mime_type == "application/pdf" or name.endswith(".pdf"):
            return _pdf(data)
        if mime_type == _DOCX or name.endswith(".docx"):
            return _docx(data)
        if mime_type == _XLSX or name.endswith(".xlsx"):
            return _xlsx(data)
        if mime_type in _TEXT_TYPES or name.endswith((".txt", ".csv", ".md")):
            return ExtractionResult("done", _clean(data.decode("utf-8", errors="replace")))
    except Exception:  # noqa: BLE001  (parsers raise many unrelated types on bad input)
        return ExtractionResult("failed")
    return ExtractionResult("unsupported")
