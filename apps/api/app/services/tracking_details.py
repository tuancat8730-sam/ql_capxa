"""Read the per-package sheets (`Goi_0x_*`) of the tracking workbook: milestones and issues.

Section V (milestones) and section VI (issues) are found by their header text in column B, so
the reader works when rows move. Read-only, like `tracking_workbook`.
"""

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DETAIL_SHEET = re.compile(r"^Goi_(\d{2})_")
MILESTONE_HEADER = "Mốc công việc"
ISSUE_HEADER = "Nội dung vướng mắc"
DATE_PATTERN = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")


@dataclass(frozen=True)
class MilestoneRow:
    package_code: str
    seq: int
    title: str
    documents: str
    timing: str
    unit: str
    result: str
    status: str
    note: str


@dataclass(frozen=True)
class IssueRow:
    package_code: str
    seq: int
    title: str
    cause: str
    unit: str
    proposal: str
    status: str
    priority: str
    basis: str


@dataclass(frozen=True)
class DetailSnapshot:
    milestones: tuple[MilestoneRow, ...]
    issues: tuple[IssueRow, ...]
    sheets_read: tuple[str, ...]


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _seq(value: Any) -> int | None:
    """STT cells hold 1.0, 2.0 ... or text; anything else ends the table."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int | float):
        return int(value)
    return int(value) if str(value).strip().isdigit() else None


def _find_header(sheet: Any, label: str) -> int | None:
    for row in range(1, sheet.max_row + 1):
        if _text(sheet[f"A{row}"].value) == "STT" and _text(sheet[f"B{row}"].value).startswith(
            label
        ):
            return row
    return None


def _table_rows(sheet: Any, header_row: int) -> list[tuple[int, list[Any]]]:
    """Data rows under a header, stopping at the first row whose STT is not a number."""
    rows: list[tuple[int, list[Any]]] = []
    row = header_row + 1
    while row <= sheet.max_row and _seq(sheet[f"A{row}"].value) is not None:
        values = [sheet[f"{column}{row}"].value for column in "ABCDEFGH"]
        rows.append((_seq(values[0]) or 0, values))
        row += 1
    return rows


def _milestone(code: str, seq: int, values: list[Any]) -> MilestoneRow:
    return MilestoneRow(
        package_code=code,
        seq=seq,
        title=_text(values[1]),
        documents=_text(values[2]),
        timing=_text(values[3]),
        unit=_text(values[4]),
        result=_text(values[5]),
        status=_text(values[6]),
        note=_text(values[7]),
    )


def _issue(code: str, seq: int, values: list[Any]) -> IssueRow:
    return IssueRow(
        package_code=code,
        seq=seq,
        title=_text(values[1]),
        cause=_text(values[2]),
        unit=_text(values[3]),
        proposal=_text(values[4]),
        status=_text(values[5]),
        priority=_text(values[6]),
        basis=_text(values[7]),
    )


def read_detail_sheets(path: Path) -> DetailSnapshot:
    """Read milestones and issues from every `Goi_0x_*` sheet; the CBDT sheet is skipped."""
    workbook = load_workbook(path, data_only=True, read_only=False)
    milestones: list[MilestoneRow] = []
    issues: list[IssueRow] = []
    sheets: list[str] = []
    for name in workbook.sheetnames:
        match = DETAIL_SHEET.match(name)
        if not match:
            continue
        code = match.group(1)
        sheet = workbook[name]
        sheets.append(name)
        if (header := _find_header(sheet, MILESTONE_HEADER)) is not None:
            milestones.extend(_milestone(code, seq, v) for seq, v in _table_rows(sheet, header))
        if (header := _find_header(sheet, ISSUE_HEADER)) is not None:
            issues.extend(_issue(code, seq, v) for seq, v in _table_rows(sheet, header))
    return DetailSnapshot(tuple(milestones), tuple(issues), tuple(sheets))


def first_date(text: str) -> str | None:
    """First dd/mm/yyyy in a text cell, as ISO (yyyy-mm-dd); None if no date is written."""
    match = DATE_PATTERN.search(text)
    if not match:
        return None
    day, month, year = match.groups()
    return f"{year}-{int(month):02d}-{int(day):02d}"


def last_date(text: str) -> str | None:
    """Last dd/mm/yyyy in a text cell, as ISO; used as the planned end of a milestone."""
    found = DATE_PATTERN.findall(text)
    if not found:
        return None
    day, month, year = found[-1]
    return f"{year}-{int(month):02d}-{int(day):02d}"
