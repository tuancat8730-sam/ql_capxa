"""Excel import of contract goods lines (SPEC 4.13): parse and validate, no database access.

The caller shows the result as a dry run and only writes when there is no error.
"""

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.services.search import fold

MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024
ONE_DONG = Decimal(1)
_HEADER_SCAN_ROWS = 10
_MIN_HEADERS = 3

# folded header text -> field
_ALIASES: dict[str, str] = {
    "stt": "line_no",
    "tt": "line_no",
    "so tt": "line_no",
    "ten hang hoa": "name",
    "ten hang": "name",
    "hang hoa": "name",
    "noi dung": "name",
    "ten": "name",
    "dvt": "unit",
    "don vi tinh": "unit",
    "don vi": "unit",
    "so luong": "quantity",
    "sl": "quantity",
    "don gia": "unit_price",
    "thanh tien": "amount",
    "bao hanh": "warranty_months",
    "bao hanh (thang)": "warranty_months",
    "thoi gian bao hanh": "warranty_months",
    "kiem dinh": "requires_calibration",
    "hieu chuan": "requires_calibration",
    "kiem dinh/hieu chuan": "requires_calibration",
    "can kiem dinh/hieu chuan": "requires_calibration",
    "xuat xu": "origin",
    "ghi chu": "notes",
}
HEADERS = (
    "STT",
    "Tên hàng hóa",
    "ĐVT",
    "Số lượng",
    "Đơn giá",
    "Thành tiền",
    "Bảo hành (tháng)",
    "Kiểm định/hiệu chuẩn",
    "Xuất xứ",
    "Ghi chú",
)
_TRUE = {"x", "co", "c", "1", "true", "yes", "y"}
_THOUSANDS = re.compile(r"^-?\d{1,3}(\.\d{3})+$")


class ImportFileError(ValueError):
    """The file is not a readable workbook with a recognisable header."""


@dataclass(frozen=True)
class RowError:
    row: int
    field: str
    message: str


@dataclass
class ParsedItem:
    row: int
    line_no: int | None
    name: str
    unit: str | None = None
    quantity: Decimal | None = None
    unit_price: Decimal | None = None
    amount: Decimal | None = None
    warranty_months: int | None = None
    requires_calibration: bool = False
    origin: str | None = None
    notes: str | None = None


@dataclass
class ParseResult:
    items: list[ParsedItem] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    rows_seen: int = 0

    @property
    def total_amount(self) -> Decimal:
        return sum((i.amount or Decimal(0) for i in self.items), Decimal(0))


def parse_number(value: Any) -> Decimal | None:
    """Numbers as Excel stores them, or Vietnamese text: `1.234.567`, `1.234,5`, `12,5`, `3 đ`."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("không phải số")
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    text = str(value).strip().replace("đ", "").replace("VND", "").replace(" ", "")
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    elif _THOUSANDS.match(text):
        text = text.replace(".", "")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("không phải số") from exc


def _text(value: Any) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def _header_map(rows: list[tuple[Any, ...]]) -> tuple[int, dict[int, str]]:
    for index, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        mapping = {
            col: _ALIASES[key]
            for col, cell in enumerate(row)
            if cell is not None and (key := fold(str(cell)).strip()) in _ALIASES
        }
        if len(set(mapping.values())) >= _MIN_HEADERS and "name" in mapping.values():
            return index, mapping
    raise ImportFileError("Không tìm thấy dòng tiêu đề (cần ít nhất cột 'Tên hàng hóa')")


def parse_workbook(content: bytes) -> ParseResult:
    if len(content) > MAX_BYTES:
        raise ImportFileError("Tệp quá lớn (tối đa 5 MB)")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)] if ws else []
    except Exception as exc:  # openpyxl raises many types for corrupt or non-xlsx input
        raise ImportFileError("Không đọc được tệp Excel (.xlsx)") from exc
    if not rows:
        raise ImportFileError("Tệp không có dữ liệu")

    head, mapping = _header_map(rows)
    result = ParseResult()
    for offset, row in enumerate(rows[head + 1 :], start=head + 2):  # 1-based sheet row
        values = {mapping[c]: row[c] for c in mapping if c < len(row)}
        if all(v is None or str(v).strip() == "" for v in values.values()):
            continue
        result.rows_seen += 1
        if result.rows_seen > MAX_ROWS:
            raise ImportFileError(f"Quá {MAX_ROWS} dòng")
        _read_row(offset, values, result)
    return result


def _number(
    result: ParseResult, row: int, name: str, label: str, values: dict[str, Any], *, minimum: int
) -> Decimal | None:
    try:
        number = parse_number(values.get(name))
    except ValueError:
        result.errors.append(RowError(row, name, f"{label} phải là số"))
        return None
    if number is not None and number < minimum:
        result.errors.append(RowError(row, name, f"{label} không được nhỏ hơn {minimum}"))
        return None
    return number


def _read_row(row: int, values: dict[str, Any], result: ParseResult) -> None:
    name = _text(values.get("name"))
    if name is None:
        result.errors.append(RowError(row, "name", "Thiếu tên hàng hóa"))
        return
    before = len(result.errors)
    quantity = _number(result, row, "quantity", "Số lượng", values, minimum=0)
    price = _number(result, row, "unit_price", "Đơn giá", values, minimum=0)
    amount = _number(result, row, "amount", "Thành tiền", values, minimum=0)
    warranty = _number(result, row, "warranty_months", "Bảo hành", values, minimum=0)
    line = _number(result, row, "line_no", "STT", values, minimum=1)
    if warranty is not None and warranty != warranty.to_integral_value():
        result.errors.append(RowError(row, "warranty_months", "Bảo hành phải là số tháng nguyên"))
    if quantity is not None and price is not None:
        expected = quantity * price
        if amount is None:
            amount = expected.quantize(ONE_DONG)
        elif abs(amount - expected) > ONE_DONG:
            result.errors.append(
                RowError(
                    row, "amount", f"Thành tiền không khớp số lượng × đơn giá ({expected:.0f})"
                )
            )
    if len(result.errors) > before:
        return
    calibration = fold(_text(values.get("requires_calibration")) or "")
    result.items.append(
        ParsedItem(
            row=row,
            line_no=int(line) if line is not None else None,
            name=name,
            unit=_text(values.get("unit")),
            quantity=quantity,
            unit_price=price,
            amount=amount,
            warranty_months=int(warranty) if warranty is not None else None,
            requires_calibration=calibration in _TRUE,
            origin=_text(values.get("origin")),
            notes=_text(values.get("notes")),
        )
    )
