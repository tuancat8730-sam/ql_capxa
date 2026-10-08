import io
from dataclasses import dataclass, field
from datetime import date
from typing import Annotated
from urllib.parse import quote

from fastapi import APIRouter, Depends, Request, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import ProjectDep, SessionDep, require
from app.core.rbac import Level
from app.core.scope import current_project_id
from app.models import Contract, Document, Package, StagePlan, User
from app.services import audit
from app.services.finance import today_local

router = APIRouter(tags=["export"])

PackageReader = Annotated[User, Depends(require("package", Level.READ))]

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LATE_FILL = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
LATE_FONT = Font(color="9C0006", bold=True)
HEAD_FILL = PatternFill("solid", start_color="D9E2F3", end_color="D9E2F3")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Milestone columns of QL-06 and the document type whose date proves each one.
MILESTONES: tuple[tuple[str, str | None], ...] = (
    ("QĐ phê duyệt E-HSMT", "etbmt_approval"),
    ("Đăng E-TBMT", "etbmt"),
    ("Đóng, mở thầu", "bid_opening_minutes"),
    ("Số nhà thầu", None),  # not recorded anywhere yet
    ("Hoàn thành đánh giá", "evaluation_report"),
    ("QĐ phê duyệt KQLCNT", "decision_selection_result"),
    ("Ký hợp đồng", "contract"),  # taken from the contract's signed date
)
STATUS_LABELS = {
    "planning": "Chuẩn bị",
    "bidding": "Đang lựa chọn nhà thầu",
    "negotiating": "Thương thảo",
    "contract_signed": "Đã ký hợp đồng",
    "executing": "Đang thực hiện",
    "accepted": "Đã nghiệm thu",
    "settled": "Đã quyết toán",
    "cancelled": "Đã hủy",
}
PACKAGE_TYPE_LABELS = {"goods": "Hàng hóa", "consulting": "Tư vấn"}


@dataclass
class Ql06Row:
    number: int
    name: str
    package_type: str
    price: int | None
    milestones: list[date | None]
    status: str
    issues: str = ""
    late: list[bool] = field(default_factory=list)


def is_late(value: date | None, planned_end: date | None, today: date) -> bool:
    """A milestone is late when it was due (S2 planned end) and is still missing, or came after."""
    if planned_end is None:
        return False
    if value is None:
        return planned_end < today
    return value > planned_end


async def collect_ql06(session: AsyncSession, today: date) -> list[Ql06Row]:
    rows: list[Ql06Row] = []
    packages = (
        (
            await session.execute(
                select(Package)
                .where(Package.deleted_at.is_(None), Package.project_id == current_project_id())
                .order_by(Package.number)
            )
        )
        .scalars()
        .all()
    )
    for package in packages:
        docs = {
            d.doc_type: d.doc_date
            for d in (
                await session.execute(
                    select(Document)
                    .where(
                        Document.package_id == package.id,
                        Document.is_current.is_(True),
                        Document.deleted_at.is_(None),
                    )
                    .order_by(Document.created_at)
                )
            )
            .scalars()
            .all()
        }
        contract = (
            (
                await session.execute(
                    select(Contract)
                    .where(Contract.package_id == package.id, Contract.deleted_at.is_(None))
                    .order_by(Contract.signed_date.nulls_last())
                )
            )
            .scalars()
            .first()
        )
        selection = (
            await session.execute(
                select(StagePlan).where(
                    StagePlan.package_id == package.id, StagePlan.stage_code == "S2_SELECTION"
                )
            )
        ).scalar_one_or_none()
        planned_end = selection.planned_end if selection else None

        values: list[date | None] = []
        for _label, doc_type in MILESTONES:
            if doc_type == "contract":
                values.append(contract.signed_date if contract else docs.get("contract"))
            elif doc_type is None:
                values.append(None)
            else:
                values.append(docs.get(doc_type))
        row = Ql06Row(
            number=package.number,
            name=package.name,
            package_type=package.package_type,
            price=int(package.package_price) if package.package_price is not None else None,
            milestones=values,
            status=STATUS_LABELS.get(package.status, package.status),
        )
        row.late = [
            doc_type is not None and is_late(v, planned_end, today)
            for (_label, doc_type), v in zip(MILESTONES, values, strict=True)
        ]
        rows.append(row)
    return rows


def build_ql06(project_name: str, rows: list[Ql06Row], today: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "QL-06"
    headers = [
        "Gói",
        "Tên gói thầu",
        "Loại",
        "Giá gói thầu (VND)",
        *(label for label, _ in MILESTONES),
        "Tình trạng",
        "Vướng mắc",
    ]
    last_col = get_column_letter(len(headers))
    ws["A1"] = "BẢNG THEO DÕI LỰA CHỌN NHÀ THẦU (QL-06)"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A2"] = f"{project_name} — xuất ngày {today:%d/%m/%Y}"
    ws.merge_cells(f"A2:{last_col}2")

    header_row = 4
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, row in enumerate(rows, start=header_row + 1):
        ws.cell(row=i, column=1, value=row.number)
        ws.cell(row=i, column=2, value=row.name)
        ws.cell(row=i, column=3, value=PACKAGE_TYPE_LABELS.get(row.package_type, row.package_type))
        price = ws.cell(row=i, column=4, value=row.price)
        price.number_format = "#,##0"
        for offset, value in enumerate(row.milestones):
            cell = ws.cell(row=i, column=5 + offset, value=value)
            cell.number_format = "DD/MM/YYYY"
            cell.alignment = Alignment(horizontal="center")
            if row.late[offset]:
                cell.fill, cell.font = LATE_FILL, LATE_FONT
        status_col = 5 + len(MILESTONES)
        ws.cell(row=i, column=status_col, value=row.status)
        ws.cell(row=i, column=status_col + 1, value=row.issues)
        for col in range(1, len(headers) + 1):
            ws.cell(row=i, column=col).border = BORDER

    widths = [6, 34, 12, 20, *[17] * len(MILESTONES), 26, 40]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[header_row].height = 34
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/export/ql06.xlsx")
async def export_ql06(
    request: Request, user: PackageReader, session: SessionDep, project: ProjectDep
) -> Response:
    today = today_local()
    rows = await collect_ql06(session, today)
    content = build_ql06(project.name, rows, today)
    audit.record(
        session,
        action="export",
        entity_type="ql06",
        entity_id=None,
        user_id=user.id,
        changes={"rows": len(rows)},
        request=request,
    )
    await session.commit()
    filename = f"QL-06-{today:%Y%m%d}.xlsx"
    disposition = f"attachment; filename=\"{filename}\"; filename*=UTF-8''{quote(filename)}"
    return Response(content=content, media_type=XLSX, headers={"Content-Disposition": disposition})
