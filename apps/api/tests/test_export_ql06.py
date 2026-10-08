"""QL-06 export (SPEC 4.4): Vietnamese headers, real dates and money, late cells in red."""

import io
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import AuditLog, Package, StagePlan
from app.routers.exports import MILESTONES, is_late
from app.seed.progress import seed_progress
from app.seed.project import seed_project
from app.services.storage import MemoryStorage
from tests.conftest import ALL_ROLES
from tests.test_documents_api import upload

TODAY = date(2026, 10, 6)


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.routers.exports.today_local", lambda: TODAY)


@pytest.fixture
async def seeded(session: AsyncSession) -> dict[int, Package]:
    await seed_project(session)
    await seed_progress(session)
    return {p.number: p for p in (await session.execute(select(Package))).scalars().all()}


async def export(c):
    resp = await c.get("/api/v1/export/ql06.xlsx")
    assert resp.status_code == 200, resp.text
    return resp, load_workbook(io.BytesIO(resp.content)).active


def test_is_late_boundaries() -> None:
    planned = date(2026, 10, 1)
    assert is_late(None, None, TODAY) is False  # no plan, nothing to be late against
    assert is_late(None, planned, TODAY) is True  # due and still missing
    assert is_late(None, date(2026, 10, 6), TODAY) is False  # due today is not yet late
    assert is_late(date(2026, 10, 1), planned, TODAY) is False  # on the planned day
    assert is_late(date(2026, 10, 2), planned, TODAY) is True  # one day after
    assert is_late(date(2026, 9, 1), planned, TODAY) is False


async def test_workbook_has_vietnamese_headers_and_one_row_per_package(
    session: AsyncSession, seeded: dict[int, Package], make_client_for
) -> None:
    c = await make_client_for(session, "viewer")
    resp, ws = await export(c)
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "QL-06-20261006.xlsx" in resp.headers["content-disposition"]
    assert ws.title == "QL-06" and ws["A1"].value == "BẢNG THEO DÕI LỰA CHỌN NHÀ THẦU (QL-06)"
    assert "xuất ngày 06/10/2026" in ws["A2"].value
    headers = [ws.cell(row=4, column=i).value for i in range(1, 14)]
    assert headers[:4] == ["Gói", "Tên gói thầu", "Loại", "Giá gói thầu (VND)"]
    assert headers[4:11] == [label for label, _ in MILESTONES]
    assert headers[11:] == ["Tình trạng", "Vướng mắc"]
    assert [ws.cell(row=r, column=1).value for r in range(5, 13)] == list(range(1, 9))
    assert ws.cell(row=13, column=1).value is None  # no stray rows
    assert ws.freeze_panes == "C5"


async def test_money_and_dates_are_real_excel_values_with_vietnamese_formats(
    session: AsyncSession, seeded: dict[int, Package], make_client_for
) -> None:
    c = await make_client_for(session, "viewer")
    _, ws = await export(c)
    price = ws["D8"]  # package 4
    assert price.value == 62_298_200_000 and price.number_format == "#,##0"
    signed = ws.cell(row=8, column=11)  # contract signed 14/09/2026
    assert signed.value.date() == date(2026, 9, 14) and signed.number_format == "DD/MM/YYYY"
    assert ws.cell(row=8, column=12).value == "Đã ký hợp đồng"
    assert ws["D7"].value == 430_000_000  # package 3: consulting, value from the listing
    assert ws.cell(row=7, column=12).value == "Đã ký hợp đồng"
    assert ws["C8"].value == "Hàng hóa" and ws["C10"].value == "Tư vấn"


async def test_dates_come_from_the_uploaded_documents(
    session: AsyncSession, seeded: dict[int, Package], storage: MemoryStorage, make_client_for
) -> None:
    c = await make_client_for(session, "clerk")
    await upload(
        c, storage, seeded[4], doc_type="etbmt", title="E-TBMT", name="t.pdf", doc_date="2026-07-10"
    )
    await upload(
        c,
        storage,
        seeded[4],
        doc_type="evaluation_report",
        title="Đánh giá",
        name="e.pdf",
        doc_date="2026-08-20",
    )
    _, ws = await export(c)
    assert ws.cell(row=8, column=6).value.date() == date(2026, 7, 10)  # Đăng E-TBMT
    assert ws.cell(row=8, column=9).value.date() == date(2026, 8, 20)  # Hoàn thành đánh giá
    assert ws.cell(row=8, column=5).value is None


async def test_late_cells_are_red_and_on_time_cells_are_not(
    session: AsyncSession, seeded: dict[int, Package], storage: MemoryStorage, make_client_for
) -> None:
    c = await make_client_for(session, "admin")
    s2 = (
        await session.execute(
            select(StagePlan).where(
                StagePlan.package_id == seeded[4].id, StagePlan.stage_code == "S2_SELECTION"
            )
        )
    ).scalar_one()
    s2.planned_end = date(2026, 8, 1)  # selection should have finished by 01/08
    await session.commit()
    await upload(
        c, storage, seeded[4], doc_type="etbmt", title="E-TBMT", name="t.pdf", doc_date="2026-07-10"
    )
    await upload(
        c,
        storage,
        seeded[4],
        doc_type="evaluation_report",
        title="Đánh giá",
        name="e.pdf",
        doc_date="2026-08-20",
    )
    _, ws = await export(c)

    def red(col: int) -> bool:
        cell = ws.cell(row=8, column=col)
        return cell.fill.start_color.rgb.endswith("FFC7CE")

    assert not red(6)  # E-TBMT on 10/07, before the 01/08 deadline
    assert red(9)  # evaluation finished 20/08, after the deadline
    assert red(5) and red(7) and red(10)  # still missing although due
    assert not red(8)  # "Số nhà thầu" has no data source: never flagged
    assert ws.cell(row=8, column=9).font.bold is True
    # a package without a plan is never flagged
    assert not any(
        ws.cell(row=5, column=i).fill.start_color.rgb.endswith("FFC7CE") for i in range(5, 12)
    )


async def test_export_is_audited(
    session: AsyncSession, seeded: dict[int, Package], make_client_for
) -> None:
    c = await make_client_for(session, "director")
    await export(c)
    row = (await session.execute(select(AuditLog).where(AuditLog.action == "export"))).scalar_one()
    assert (row.entity_type, row.changes, row.user_id is not None) == ("ql06", {"rows": 8}, True)


@pytest.mark.parametrize("role", ALL_ROLES)
async def test_any_role_that_can_read_packages_can_export(
    role: str, session: AsyncSession, seeded: dict[int, Package], make_client_for
) -> None:
    c = await make_client_for(session, role)
    assert (await c.get("/api/v1/export/ql06.xlsx")).status_code == 200


async def test_export_requires_login(client) -> None:
    assert (await client.get("/api/v1/export/ql06.xlsx")).status_code == 401
