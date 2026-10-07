"""A synthetic delivery-plan .docx with the same layout as the contractors' documents."""

import io
from datetime import date

import docx

HEADER_LINES = [
    "KẾ HOẠCH TRIỂN KHAI",
    "Kính gửi: Sở Khoa học và Công nghệ tỉnh Lâm Đồng",
    "Căn cứ Hợp đồng số 99/2026/HĐ ký ngày 14/9/2026 giữa Sở và Liên danh Gói thầu số 04.",
    "Thời gian thực hiện hợp đồng: 60 ngày, kể từ ngày 14/9/2026 đến ngày 13/11/2026.",
    "Thời gian triển khai: Từ ngày 05/11/2026 đến ngày 13/11/2026.",
    "Địa điểm triển khai:",
    "- Kiểm tra đầu vào: tại kho của Sở (36 Trần Phú, Đà Lạt).",
    "- Giao hàng và lắp đặt: tại trụ sở 87 UBND xã, phường.",
    "4. Nội dung triển khai:",
]

ITEMS_HEADER = ["Stt", "Tên thiết bị", "ĐVT", "Khối\nlượng", "Nhà thầu đảm nhiệm", "Ghi chú"]
STEPS_HEADER = ["STT", "Nội dung công việc", "Thời gian", "Thành phần tham gia"]


def _fill(cell, text: str) -> None:
    lines = text.split("\n")
    cell.paragraphs[0].text = lines[0]
    for extra in lines[1:]:
        cell.add_paragraph(extra)


def build_plan_docx(
    *,
    items: list[list[str]] | None = None,
    steps: list[tuple[str, list[str]]] | None = None,
    total: str = "9",
    header: list[str] | None = None,
) -> bytes:
    """`steps` is a list of (group heading or "", row) pairs; a heading starts a new stage row."""
    document = docx.Document()
    for line in header if header is not None else HEADER_LINES:
        document.add_paragraph(line)
    items = (
        items
        if items is not None
        else [
            [
                "1",
                "Cân phân tích 220 g\nHãng X – Model Y\n- Bảo hành: 12 tháng",
                "Cái",
                "6",
                "CÔNG TY TNHH A",
                "",
            ],
            [
                "2",
                "Bộ phần mềm kèm máy tính",
                "Bộ",
                "3",
                "Công ty TNHH A: 01 bộ\nCông ty TNHH B: 02 bộ",
                "Hàng nhập khẩu",
            ],
        ]
    )
    table = document.add_table(rows=1, cols=6)
    for cell, text in zip(table.rows[0].cells, ITEMS_HEADER, strict=True):
        _fill(cell, text)
    for row in items:
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            _fill(cell, text)
    cells = table.add_row().cells
    _fill(cells[1], "TỔNG CỘNG")
    _fill(cells[3], total)

    document.add_paragraph("4.2 Các bước triển khai:")
    steps = (
        steps
        if steps is not None
        else [
            (
                "BƯỚC 1: KIỂM TRA HÀNG HÓA TẠI KHO",
                [
                    "1",
                    "- Kiểm tra hàng tại kho.\n- Lập biên bản.",
                    "Từ ngày 30/10/2026 đến ngày 05/11/2026",
                    "Nhà thầu",
                ],
            ),
            (
                "",
                [
                    "2",
                    "- Gửi thiết bị đi kiểm định.",
                    "Dự trù từ ngày 06/11/2026 đến ngày 15/11/2026\n"
                    "(Tùy thuộc vào tổ chức kiểm định)",
                    "Tổ chức kiểm định (PA05, PA06), Nhà thầu",
                ],
            ),
            (
                "BƯỚC 2: LẮP ĐẶT, NGHIỆM THU",
                ["3", "- Thông báo kết quả.", "Dự kiến ngày 15/11/2026", "Sở, Đơn vị QLDA"],
            ),
            ("", ["4", "- Nghiệm thu.", "27/11/2026", "Sở, Nhà thầu"]),
            ("", ["5", "- Bàn giao hồ sơ.", "", "Sở"]),
        ]
    )
    steps_table = document.add_table(rows=1, cols=4)
    for cell, text in zip(steps_table.rows[0].cells, STEPS_HEADER, strict=True):
        _fill(cell, text)
    for heading, row in steps:
        if heading:
            merged = steps_table.add_row().cells
            merged[0].merge(merged[3])
            _fill(merged[0], heading)
        cells = steps_table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            _fill(cell, text)
    document.add_paragraph("TM. LIÊN DANH GÓI THẦU SỐ 04")
    document.add_paragraph("THÀNH VIÊN ĐỨNG ĐẦU LIÊN DANH")
    document.add_paragraph("CÔNG TY TNHH A")
    document.add_paragraph("Giám đốc")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


CONTRACT_END = date(2026, 11, 13)
