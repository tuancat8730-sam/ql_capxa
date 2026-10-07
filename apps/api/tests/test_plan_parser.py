"""Reading a contractor's delivery plan: layout, dates, assignments and consistency findings."""

import io
import os
import zipfile
from datetime import date
from pathlib import Path

import pytest

from app.services.plan_parser import (
    Finding,
    ParsedPlan,
    content_hash,
    parse_dates,
    parse_plan,
    plan_findings,
)
from app.services.plan_reader import MAX_BYTES, PlanFileError, read_plan_document
from tests.plan_docs import build_plan_docx

D = date


def parse(**kw) -> ParsedPlan:
    return parse_plan(read_plan_document(build_plan_docx(**kw), "plan.docx"))


def codes(findings: list[Finding]) -> list[str]:
    return [f.code for f in findings]


# --- header --------------------------------------------------------------------------------------


def test_header_facts() -> None:
    plan = parse()
    assert plan.addressee == "Sở Khoa học và Công nghệ tỉnh Lâm Đồng"
    assert plan.legal_basis and plan.legal_basis.startswith("Căn cứ Hợp đồng số 99/2026")
    assert (plan.contract_start, plan.contract_end) == (D(2026, 9, 14), D(2026, 11, 13))
    assert (plan.implement_start, plan.implement_end) == (D(2026, 11, 5), D(2026, 11, 13))
    assert plan.contract_text and plan.contract_text.startswith("60 ngày")
    assert plan.locations == [
        {"label": "Kiểm tra đầu vào", "text": "tại kho của Sở (36 Trần Phú, Đà Lạt)."},
        {"label": "Giao hàng và lắp đặt", "text": "tại trụ sở 87 UBND xã, phường."},
    ]
    assert plan.signer == "CÔNG TY TNHH A – Giám đốc"


def test_missing_header_facts_stay_empty() -> None:
    plan = parse(header=["Một văn bản chỉ có hai bảng"])
    assert plan.addressee is None and plan.contract_end is None and plan.locations == []
    assert plan.items and plan.steps


# --- equipment -----------------------------------------------------------------------------------


def test_items_split_name_details_quantity_and_assignments() -> None:
    plan = parse()
    first, second = plan.items
    assert (first.line_no, first.name, first.unit, first.quantity) == (
        1,
        "Cân phân tích 220 g",
        "Cái",
        6,
    )
    assert first.details == "Hãng X – Model Y\n- Bảo hành: 12 tháng"
    # one contractor does the whole line
    assert first.assignments == [{"org": "CÔNG TY TNHH A", "quantity": 6}]
    assert second.assignments == [
        {"org": "Công ty TNHH A", "quantity": 1},
        {"org": "Công ty TNHH B", "quantity": 2},
    ]
    assert second.note == "Hàng nhập khẩu"
    assert plan.declared_total == 9 and sum(i.quantity or 0 for i in plan.items) == 9


def test_quantities_with_thousand_separators() -> None:
    plan = parse(items=[["1", "Cáp", "m", "1.250", "CÔNG TY A", ""]], total="1.250")
    assert plan.items[0].quantity == 1250 and plan.declared_total == 1250


# --- steps ---------------------------------------------------------------------------------------


def test_steps_carry_their_stage_and_dates() -> None:
    steps = parse().steps
    assert [(s.group_no, s.step_no) for s in steps] == [(1, 1), (1, 2), (2, 3), (2, 4), (2, 5)]
    assert steps[0].group_title == "KIỂM TRA HÀNG HÓA TẠI KHO"
    assert (steps[0].start_date, steps[0].end_date, steps[0].estimated) == (
        D(2026, 10, 30),
        D(2026, 11, 5),
        False,
    )
    assert steps[0].content == "- Kiểm tra hàng tại kho.\n- Lập biên bản."
    assert steps[0].participants == ["Nhà thầu"]


def test_estimated_dates_and_bracket_notes() -> None:
    estimated = parse().steps[1]
    assert estimated.estimated is True
    assert (estimated.start_date, estimated.end_date) == (D(2026, 11, 6), D(2026, 11, 15))
    assert estimated.time_note == "Tùy thuộc vào tổ chức kiểm định"
    # commas inside brackets do not split a participant
    assert estimated.participants == ["Tổ chức kiểm định (PA05, PA06)", "Nhà thầu"]
    single = parse().steps[2]
    assert (single.start_date, single.end_date, single.estimated) == (
        D(2026, 11, 15),
        D(2026, 11, 15),
        True,
    )


def test_a_single_plain_date_is_a_one_day_step() -> None:
    step = parse().steps[3]
    assert (step.start_date, step.end_date, step.estimated) == (
        D(2026, 11, 27),
        D(2026, 11, 27),
        False,
    )


def test_a_step_without_time_has_no_dates() -> None:
    last = parse().steps[4]
    assert (last.time_text, last.start_date, last.end_date) == (None, None, None)


def test_impossible_dates_are_reported_not_guessed() -> None:
    plan = parse(
        steps=[
            (
                "BƯỚC 1: A",
                ["1", "- Làm việc.", "Từ ngày 31/02/2026 đến ngày 05/03/2026", "Nhà thầu"],
            )
        ]
    )
    assert plan.steps[0].start_date == D(2026, 3, 5)
    assert any("31/02/2026" in w for w in plan.warnings)
    assert parse_dates("12/10/2026 và 31/04/2026") == ([D(2026, 10, 12)], ["31/04/2026"])


def test_steps_before_any_stage_heading_have_no_stage() -> None:
    plan = parse(steps=[("", ["1", "- Việc đầu tiên.", "01/12/2026", "Nhà thầu"])])
    assert (plan.steps[0].group_no, plan.steps[0].group_title) == (None, None)


# --- findings ------------------------------------------------------------------------------------


def test_findings_flag_what_does_not_add_up() -> None:
    found = {f.code: f for f in plan_findings(parse())}
    assert found["AFTER_CONTRACT_END"].severity == "warning"
    assert (
        "13/11/2026" in found["AFTER_CONTRACT_END"].message
        and "27/11/2026" in found["AFTER_CONTRACT_END"].message
    )
    assert found["AFTER_IMPLEMENT_END"].severity == "info"
    assert found["NO_TIME"].message.endswith("bước 5")
    assert "QUANTITY_TOTAL" not in found


def test_a_plan_inside_the_contract_period_is_clean() -> None:
    plan = parse(
        steps=[("BƯỚC 1: A", ["1", "- Làm.", "Từ ngày 05/11/2026 đến ngày 12/11/2026", "Nhà thầu"])]
    )
    assert plan_findings(plan) == []


def test_quantity_findings() -> None:
    wrong_total = parse(total="10")
    assert "QUANTITY_TOTAL" in codes(plan_findings(wrong_total))
    split = parse(items=[["1", "Máy", "Bộ", "10", "A: 3 bộ\nB: 4 bộ", ""]], total="10")
    assert "ASSIGNMENT_TOTAL" in codes(plan_findings(split))


def test_contract_end_of_the_package_is_used_when_the_plan_has_none() -> None:
    plan = parse(header=["Một văn bản không nêu thời gian hợp đồng"])
    assert "AFTER_CONTRACT_END" not in codes(plan_findings(plan))
    assert "AFTER_CONTRACT_END" in codes(plan_findings(plan, contract_end=D(2026, 11, 1)))


def test_content_hash_ignores_case_accents_and_spacing() -> None:
    assert content_hash("- Nghiệm thu   lắp đặt") == content_hash("- nghiem thu lap dat")
    assert content_hash("- Nghiệm thu") != content_hash("- Bàn giao")


# --- refusing bad files --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("data", "name", "message"),
    [
        (b"hello", "plan.docx", "không phải văn bản Word"),
        (b"PK\x03\x04junk", "plan.docx", "Không đọc được"),
        (build_plan_docx(), "plan.pdf", "Chỉ nhận tệp Word"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 50, "plan.doc", ".docx"),
    ],
)
def test_bad_files_are_refused_with_a_readable_message(
    data: bytes, name: str, message: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("ANTIWORD_BIN", "/nonexistent/antiword")
    with pytest.raises(PlanFileError, match=message):
        read_plan_document(data, name)


def test_oversized_and_bombed_files_are_refused() -> None:
    with pytest.raises(PlanFileError, match="quá lớn"):
        read_plan_document(b"0" * (MAX_BYTES + 1), "plan.docx")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<w/>")
        z.writestr("big.bin", b"0" * (60 * 1024 * 1024))  # compresses to almost nothing
    with pytest.raises(PlanFileError, match="bất thường"):
        read_plan_document(buf.getvalue(), "plan.docx")


def test_macro_documents_are_refused() -> None:
    buf = io.BytesIO(build_plan_docx())
    with zipfile.ZipFile(buf, "a") as z:
        z.writestr("word/vbaProject.bin", b"x")
    with pytest.raises(PlanFileError, match="macro"):
        read_plan_document(buf.getvalue(), "plan.docx")


def test_a_word_file_without_tables_is_refused() -> None:
    import docx

    document = docx.Document()
    document.add_paragraph("Chỉ có chữ")
    out = io.BytesIO()
    document.save(out)
    with pytest.raises(PlanFileError, match="bảng"):
        read_plan_document(out.getvalue(), "plan.docx")


def test_tables_that_are_not_a_plan_are_refused() -> None:
    import docx

    document = docx.Document()
    t = document.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Họ tên"
    t.rows[0].cells[1].text = "Tuổi"
    out = io.BytesIO()
    document.save(out)
    with pytest.raises(PlanFileError, match="Không nhận ra"):
        parse_plan(read_plan_document(out.getvalue(), "x.docx"))


# --- the real documents (kept out of git: run only where the files exist) ----------------------

DATA = Path(__file__).resolve().parents[3] / "data_documents"
GOI4 = DATA / "Ke_hoach_trien_khai_lap_dat_Goi04_HC edit 9h50 07.10.docx"
GOI5 = DATA / "Kế hoạch triển khai_lap_dat_GOI_5.doc"


@pytest.mark.skipif(not GOI4.exists(), reason="real plan documents are not in the repository")
def test_real_package_4_plan() -> None:
    plan = parse_plan(read_plan_document(GOI4.read_bytes(), GOI4.name))
    assert (
        len(plan.items) == 10
        and sum(i.quantity or 0 for i in plan.items) == 846 == plan.declared_total
    )
    assert [s.step_no for s in plan.steps] == list(range(1, 11))
    assert {s.group_no for s in plan.steps} == {1, 2, 3}
    assert plan.steps[-1].start_date is None  # step 10 has no time
    assert "AFTER_CONTRACT_END" in codes(plan_findings(plan))


@pytest.mark.skipif(
    not GOI5.exists()
    or not (os.environ.get("ANTIWORD_BIN") or __import__("shutil").which("antiword")),
    reason="needs the real .doc file and antiword",
)
def test_real_package_5_plan_from_an_old_doc_file() -> None:
    plan = parse_plan(read_plan_document(GOI5.read_bytes(), GOI5.name))
    assert len(plan.items) == 9 and sum(i.quantity or 0 for i in plan.items) == 751
    assert [s.step_no for s in plan.steps] == list(range(1, 12))
    assert (plan.implement_start, plan.implement_end) == (D(2026, 10, 12), D(2026, 11, 13))
    assert plan_findings(plan) == []  # every step is inside the contract period
