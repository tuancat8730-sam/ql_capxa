import io

from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from app.services.extract import extract_text

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_text_pdf(text: str) -> bytes:
    """Smallest valid one-page PDF carrying `text` (ASCII, Helvetica)."""
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def test_pdf_with_text_layer_is_extracted() -> None:
    res = extract_text(make_text_pdf("Bao lanh tam ung so 123 cua ngan hang"), PDF, "a.pdf")
    assert res.status == "done"
    assert res.text is not None and "Bao lanh tam ung so 123" in res.text


def test_scanned_pdf_is_flagged_for_ocr_not_failed() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=144)
    buf = io.BytesIO()
    writer.write(buf)
    res = extract_text(buf.getvalue(), PDF, "scan.pdf")
    assert (res.status, res.text) == ("needs_ocr", None)


def test_pdf_with_only_a_few_characters_counts_as_scan() -> None:
    assert extract_text(make_text_pdf("p. 1"), PDF).status == "needs_ocr"


def test_docx_paragraphs_and_tables_are_extracted_with_diacritics() -> None:
    doc = DocxDocument()
    doc.add_paragraph("Bảo lãnh tạm ứng")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Nhà thầu"
    table.rows[0].cells[1].text = "Nguyên Luân"
    buf = io.BytesIO()
    doc.save(buf)
    res = extract_text(buf.getvalue(), DOCX, "x.docx")
    assert res.status == "done" and res.text is not None
    assert "Bảo lãnh tạm ứng" in res.text and "Nguyên Luân" in res.text


def test_xlsx_cells_and_sheet_names_are_extracted() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hàng hóa"
    ws.append(["Máy tính", 12])
    buf = io.BytesIO()
    wb.save(buf)
    res = extract_text(buf.getvalue(), XLSX, "x.xlsx")
    assert res.status == "done" and res.text is not None
    assert "Hàng hóa" in res.text and "Máy tính" in res.text and "12" in res.text


def test_plain_text_is_decoded_and_whitespace_collapsed() -> None:
    res = extract_text("Biên   bản\n\nnghiệm thu".encode(), "text/plain", "n.txt")
    assert res.text == "Biên bản nghiệm thu"


def test_nul_bytes_are_removed() -> None:
    assert extract_text(b"a\x00b", "text/plain").text == "a b"


def test_images_zip_and_legacy_office_are_unsupported() -> None:
    for mime, name in (
        ("image/jpeg", "a.jpg"),
        ("application/zip", "a.zip"),
        ("application/msword", "a.doc"),
    ):
        assert extract_text(b"data", mime, name).status == "unsupported"


def test_corrupt_files_become_failed_instead_of_raising() -> None:
    assert extract_text(b"not a pdf", PDF, "a.pdf").status == "failed"
    assert extract_text(b"not a zip", DOCX, "a.docx").status == "failed"
    assert extract_text(b"not a zip", XLSX, "a.xlsx").status == "failed"


def test_type_is_detected_from_the_file_name_when_mime_is_generic() -> None:
    res = extract_text(
        make_text_pdf("Bien ban nghiem thu thiet bi"), "application/octet-stream", "n.PDF"
    )
    assert res.status == "done"


def test_text_is_capped() -> None:
    res = extract_text((("x" * 10 + " ") * 400_000 + "tail").encode(), "text/plain")
    assert res.text is not None and len(res.text) <= 2_000_000
