"""M8: Excel list exports (SPEC 4.13) - Vietnamese headers, real dates and money, audited."""

import io
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    AuditLog,
    ChecklistItem,
    Contract,
    Document,
    Issue,
    Package,
    Payment,
    Project,
)
from app.seed.checklists import seed_checklists
from app.seed.finance import seed_finance
from app.seed.progress import seed_progress
from app.seed.project import seed_project
from app.seed.risks import seed_risks
from app.services.alert_engine import run_alerts
from tests.conftest import ALL_ROLES, make_user

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NOW = datetime(2026, 11, 5, 3, 0, tzinfo=UTC)
HEAD = 4  # worksheet row holding the column headers


@pytest.fixture
async def seeded(session: AsyncSession) -> None:
    await seed_project(session)
    await seed_finance(session)
    await seed_checklists(session)
    await seed_progress(session)
    await seed_risks(session)
    await run_alerts(session, NOW)


def sheet(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    return load_workbook(io.BytesIO(response.content)).active


def table(ws) -> list[list]:
    return [[c.value for c in row] for row in ws.iter_rows(min_row=HEAD + 1)]


# --- QL-07 -------------------------------------------------------------------------------------


async def test_ql07_has_one_row_per_contract_and_a_totals_row(
    session: AsyncSession, seeded: None, make_client_for
) -> None:
    c = await make_client_for(session, "cost")
    r = await c.get("/api/v1/export/ql07.xlsx")
    ws = sheet(r)
    assert ws["A1"].value == "BẢNG THEO DÕI HỢP ĐỒNG (QL-07)"
    headers = [cell.value for cell in ws[HEAD]]
    assert headers[:5] == ["Gói", "Nhà thầu", "Số hợp đồng", "Ngày ký", "Giá trị hợp đồng (VND)"]
    assert len(headers) == 14 and headers[-1] == "Tạm ứng còn dư (VND)"
    rows = table(ws)
    contracts = (await session.execute(select(Contract))).scalars().all()
    assert len(rows) == len(contracts) + 1 and rows[-1][0] == "Tổng cộng"
    assert [row[0] for row in rows[:-1]] == sorted(row[0] for row in rows[:-1])  # by package
    assert 'filename="QL-07-' in r.headers["content-disposition"]


async def test_ql07_columns_match_the_data(
    session: AsyncSession, seeded: None, make_client_for
) -> None:
    six = (await session.execute(select(Package).where(Package.number == 6))).scalar_one()
    contract = (
        await session.execute(select(Contract).where(Contract.package_id == six.id))
    ).scalar_one()
    advance = (
        (
            await session.execute(
                select(Payment).where(
                    Payment.contract_id == contract.id, Payment.payment_type == "advance"
                )
            )
        )
        .scalars()
        .first()
    )
    assert advance
    advance.status, advance.paid_date = "paid", date(2026, 10, 1)
    session.add(
        Payment(
            contract_id=contract.id,
            payment_type="recovery",
            seq=1,
            amount=Decimal(100_000_000),
            status="paid",
            paid_date=date(2026, 10, 20),
        )
    )
    await session.commit()

    c = await make_client_for(session, "director")
    rows = table(sheet(await c.get("/api/v1/export/ql07.xlsx")))
    row = next(r for r in rows if r[0] == "Gói 06")
    assert row[1] == "Sài Gòn Mới" and row[2] == "80"
    assert row[3].date() == date(2026, 9, 24)
    assert row[4] == 1_726_920_000 and row[8] == 518_076_000
    assert row[11] == 518_076_000  # cumulative paid
    assert row[12] == 100_000_000  # recovered
    assert row[13] == 418_076_000  # advance 518.076.000 - 100.000.000
    total = rows[-1]
    assert total[4] == sum(r[4] or 0 for r in rows[:-1])
    assert total[13] == sum(r[13] or 0 for r in rows[:-1])


async def test_ql07_lists_the_guarantee_in_force(
    session: AsyncSession, seeded: None, make_client_for
) -> None:
    c = await make_client_for(session, "director")
    rows = table(sheet(await c.get("/api/v1/export/ql07.xlsx")))
    five = next(r for r in rows if r[0] == "Gói 05")
    assert five[9] is not None and five[10].date() == date(2026, 11, 13)  # advance guarantee
    four = next(r for r in rows if r[0] == "Gói 04")
    assert four[9] is None  # its guarantee is only "missing": nothing to list


# --- risks, issues, alerts, checklist ----------------------------------------------------------


async def test_risk_export_orders_by_score_and_marks_high_risks(
    session: AsyncSession, seeded: None, make_client_for
) -> None:
    c = await make_client_for(session, "technical")
    ws = sheet(await c.get("/api/v1/export/risks.xlsx"))
    assert ws["A1"].value == "SỔ RỦI RO"
    rows = table(ws)
    assert len(rows) == 10
    scores = [r[6] for r in rows]
    assert scores == sorted(scores, reverse=True) and rows[0][7] == "Cao"
    assert all(r[6] == r[4] * r[5] for r in rows)
    assert ws.cell(row=HEAD + 1, column=7).fill.start_color.rgb.endswith("FFC7CE")
    low = next(i for i, r in enumerate(rows) if r[7] != "Cao")
    assert ws.cell(row=HEAD + 1 + low, column=7).fill.fill_type is None


async def test_issue_export_flags_overdue_deadlines(
    session: AsyncSession, seeded: None, make_client_for
) -> None:
    reporter = await make_user(session, "technical")
    overdue = Issue(
        code="V-001",
        issue_type="contract",
        level=1,
        title="Quá hạn",
        reported_by=reporter.id,
        due_at=datetime.now(UTC) - timedelta(days=2),
        status="open",
    )
    fine = Issue(
        code="V-002",
        issue_type="operational",
        level=1,
        title="Còn hạn",
        reported_by=reporter.id,
        due_at=datetime.now(UTC) + timedelta(days=2),
        status="open",
    )
    done = Issue(
        code="V-003",
        issue_type="other",
        level=2,
        title="Xong",
        reported_by=reporter.id,
        due_at=datetime.now(UTC) - timedelta(days=9),
        status="resolved",
    )
    session.add_all([overdue, fine, done])
    await session.commit()

    c = await make_client_for(session, "director")
    ws = sheet(await c.get("/api/v1/export/issues.xlsx"))
    rows = table(ws)
    assert [r[0] for r in rows] == ["V-001", "V-002", "V-003"]
    assert rows[0][2] == "Hợp đồng" and rows[0][5] == "Mới"
    hot = [ws.cell(row=HEAD + 1 + i, column=8).fill.fill_type for i in range(3)]
    assert hot == ["solid", None, None]  # only the open overdue one is red


async def test_alert_export_lists_critical_first_without_resolved_ones(
    session: AsyncSession, seeded: None, make_client_for
) -> None:
    c = await make_client_for(session, "viewer")
    ws = sheet(await c.get("/api/v1/export/alerts.xlsx"))
    rows = table(ws)
    assert rows and rows[0][0] == "Nghiêm trọng"
    order = {"Nghiêm trọng": 0, "Cần chú ý": 1, "Thông tin": 2}
    ranks = [order[r[0]] for r in rows]
    assert ranks == sorted(ranks)
    assert all(r[6] != "Đã tự đóng" for r in rows)
    assert any(r[1] == "Bảo lãnh tạm ứng không đủ thời hạn" and r[2] == "Gói 05" for r in rows)


async def test_checklist_export_has_every_item_and_hides_hidden_files(
    session: AsyncSession, seeded: None, make_client_for
) -> None:
    item = (
        (await session.execute(select(ChecklistItem).where(ChecklistItem.status == "missing")))
        .scalars()
        .first()
    )
    assert item
    project = (await session.execute(select(Project))).scalars().one()
    doc = Document(
        project_id=project.id,
        package_id=item.package_id,
        category="contract",
        doc_type=item.doc_type,
        title="Mật",
        confidentiality="sensitive",
        file_key="k",
        file_name="m.pdf",
        mime_type="application/pdf",
        size_bytes=1,
        sha256="0" * 64,
        doc_date=date(2026, 8, 8),
    )
    session.add(doc)
    await session.flush()
    item.status, item.document_id = "received", doc.id
    await session.commit()

    total = len((await session.execute(select(ChecklistItem))).scalars().all())
    director = await make_client_for(session, "director")
    viewer = await make_client_for(session, "viewer")
    seen = table(sheet(await director.get("/api/v1/export/checklist.xlsx")))
    hidden = table(sheet(await viewer.get("/api/v1/export/checklist.xlsx")))
    assert len(seen) == len(hidden) == total
    assert any(r[6] is not None and r[4] == "Đã có" for r in seen)
    assert all(r[6] is None for r in hidden)  # the date of a sensitive file stays private


# --- access and audit --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("path", "allowed"),
    [
        ("ql07", {"admin", "director", "procurement", "technical", "cost", "viewer"}),
        ("risks", {"admin", "director", "procurement", "technical", "cost", "onsite", "viewer"}),
        ("issues", {"admin", "director", "procurement", "technical", "cost", "onsite", "viewer"}),
        ("alerts", set(ALL_ROLES)),
        ("checklist", set(ALL_ROLES)),
    ],
)
async def test_role_matrix(
    path: str, allowed: set[str], session: AsyncSession, seeded: None, make_client_for, client
) -> None:
    assert (await client.get(f"/api/v1/export/{path}.xlsx")).status_code == 401
    for role in ALL_ROLES:
        c = await make_client_for(session, role)
        status = (await c.get(f"/api/v1/export/{path}.xlsx")).status_code
        assert status == (200 if role in allowed else 403), (path, role)


async def test_every_export_is_audited(
    session: AsyncSession, seeded: None, make_client_for
) -> None:
    c = await make_client_for(session, "director")
    for path in ("ql07", "risks", "issues", "alerts", "checklist"):
        assert (await c.get(f"/api/v1/export/{path}.xlsx")).status_code == 200
    rows = (
        (await session.execute(select(AuditLog).where(AuditLog.action == "export"))).scalars().all()
    )
    assert {r.entity_type for r in rows} == {
        "QL-07",
        "rui-ro",
        "vuong-mac",
        "canh-bao",
        "danh-muc-ho-so",
    }
    assert all(r.user_id is not None and "rows" in r.changes for r in rows)


async def test_exports_without_data_still_open(session: AsyncSession, make_client_for) -> None:
    await seed_project(session)
    c = await make_client_for(session, "director")
    assert table(sheet(await c.get("/api/v1/export/risks.xlsx"))) == []
    assert table(sheet(await c.get("/api/v1/export/alerts.xlsx"))) == []
