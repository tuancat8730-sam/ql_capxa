"""Delivery-plan steps in the alert engine and on the dashboard (PLAN_STEP_OVERDUE)."""

from datetime import UTC, date, datetime, timedelta

import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Alert, Package, PlanStep, User
from app.seed.project import seed_project
from app.services import alert_rules as rules
from app.services.alert_engine import run_alerts
from app.services.plan_parser import parse_plan
from app.services.plan_reader import read_plan_document
from app.services.plans import save_plan
from tests.conftest import make_user
from tests.plan_docs import build_plan_docx

TODAY = date(2026, 11, 20)
NOW = datetime(2026, 11, 20, 3, 0, tzinfo=UTC)  # 10:00 in Vietnam


# --- the rule ------------------------------------------------------------------------------------


def _rule(end: date | None, status: str = "not_started", content: str = "- Nghiệm thu lắp đặt."):
    return rules.plan_step_overdue(
        step_id="s1",
        package_id="p4",
        package_label="Gói 04",
        step_no=8,
        content=content,
        end_date=end,
        status=status,
        today=TODAY,
    )


@pytest.mark.parametrize(
    ("days_late", "expected"),
    [
        (-3, []),
        (0, []),  # due today: not yet late
        (1, [("PLAN_STEP_OVERDUE", "warning")]),
        (5, [("PLAN_STEP_OVERDUE", "warning")]),
        (6, [("PLAN_STEP_OVERDUE", "critical")]),
    ],
)
def test_overdue_boundaries(days_late: int, expected: list[tuple[str, str]]) -> None:
    out = _rule(TODAY - timedelta(days=days_late))
    assert [(c.alert_type, c.severity) for c in out] == expected


def test_done_steps_and_steps_without_a_date_never_alert() -> None:
    assert _rule(TODAY - timedelta(days=30), status="done") == []
    assert _rule(None) == []
    assert len(_rule(TODAY - timedelta(days=2), status="blocked")) == 1  # blocked is still open


def test_the_alert_names_the_step_and_how_late_it_is() -> None:
    (alert,) = _rule(
        date(2026, 11, 15), content="- Nghiệm thu lắp đặt thiết bị.\n- Nghiệm thu phần mềm."
    )
    assert "bước 8" in alert.title and alert.title.startswith("Gói 04")
    assert alert.message == "Nghiệm thu lắp đặt thiết bị. - quá hạn 5 ngày (hạn 15/11/2026)"
    assert alert.due_date == date(2026, 11, 15) and alert.package_id == "p4"
    assert alert.fingerprint == "PLAN_STEP_OVERDUE:s1:"


def test_a_long_first_line_is_shortened() -> None:
    (alert,) = _rule(date(2026, 11, 1), content="- " + "rất dài " * 40)
    assert alert.message.split(" - quá hạn")[0].endswith("...") and len(alert.message) < 160


# --- the engine ----------------------------------------------------------------------------------


@pytest.fixture
async def plan_for_package_4(session: AsyncSession) -> Package:
    await seed_project(session)
    package = (await session.execute(select(Package).where(Package.number == 4))).scalar_one()
    user = await make_user(session, "admin")
    await _save(
        session, package, user, _steps(("01/11/2026", "05/11/2026"), ("10/11/2026", "18/11/2026"))
    )
    return package


def _steps(*periods: tuple[str, str]) -> list[tuple[str, list[str]]]:
    rows: list[tuple[str, list[str]]] = []
    for n, (a, b) in enumerate(periods, start=1):
        heading = "BƯỚC 1: THỰC HIỆN" if n == 1 else ""
        rows.append(
            (heading, [str(n), f"- Công việc số {n}.", f"Từ ngày {a} đến ngày {b}", "Nhà thầu"])
        )
    return rows


async def _save(session: AsyncSession, package: Package, user: User, steps) -> None:
    parsed = parse_plan(read_plan_document(build_plan_docx(steps=steps), "p.docx"))
    await save_plan(session, package, parsed, file_name="p.docx", user_id=user.id)
    await session.commit()


async def plan_alerts(session: AsyncSession) -> list[Alert]:
    stmt = (
        select(Alert)
        .where(Alert.alert_type == "PLAN_STEP_OVERDUE")
        .order_by(Alert.due_date)
        .execution_options(populate_existing=True)
    )
    return list((await session.execute(stmt)).scalars().all())


async def test_the_engine_raises_one_alert_per_overdue_step(
    session: AsyncSession, plan_for_package_4: Package
) -> None:
    result = await run_alerts(session, NOW)
    alerts = await plan_alerts(session)
    # step 1 ended 05/11 and step 2 ended 18/11: both are open on 20/11
    assert len(alerts) == 2 and result.created >= 2
    assert {a.fingerprint.split(":")[0] for a in alerts} == {"PLAN_STEP_OVERDUE"}


async def test_severity_follows_how_late_the_step_is(
    session: AsyncSession, plan_for_package_4: Package
) -> None:
    await run_alerts(session, NOW)
    alerts = await plan_alerts(session)
    by_due = {a.due_date: a for a in alerts}
    assert by_due[date(2026, 11, 5)].severity == "critical"  # 15 days late
    assert by_due[date(2026, 11, 18)].severity == "warning"  # 2 days late
    assert all(
        a.package_id == plan_for_package_4.id and a.entity_type == "plan_step" for a in alerts
    )
    assert all(a.title.startswith("Gói 04: bước") for a in alerts)


async def test_marking_the_step_done_closes_its_alert(
    session: AsyncSession, plan_for_package_4: Package
) -> None:
    await run_alerts(session, NOW)
    step = (await session.execute(select(PlanStep).where(PlanStep.step_no == 1))).scalar_one()
    step.status = "done"
    await session.commit()
    result = await run_alerts(session, NOW)
    assert result.resolved == 1
    alerts = {a.due_date: a.status for a in await plan_alerts(session)}
    assert alerts[date(2026, 11, 5)] == "resolved" and alerts[date(2026, 11, 18)] == "open"


async def test_running_again_creates_no_duplicates(
    session: AsyncSession, plan_for_package_4: Package
) -> None:
    await run_alerts(session, NOW)
    before = [a.id for a in await plan_alerts(session)]
    again = await run_alerts(session, NOW)
    assert (again.created, again.resolved) == (0, 0)
    assert [a.id for a in await plan_alerts(session)] == before


async def test_replacing_the_plan_closes_alerts_of_steps_that_no_longer_exist(
    session: AsyncSession, plan_for_package_4: Package
) -> None:
    await run_alerts(session, NOW)
    admin = (await session.execute(select(User))).scalars().first()
    assert admin
    await _save(session, plan_for_package_4, admin, _steps(("01/12/2026", "05/12/2026")))
    await run_alerts(session, NOW)
    assert {a.status for a in await plan_alerts(session)} == {"resolved"}


async def test_critical_plan_alert_makes_the_package_red(
    session: AsyncSession, plan_for_package_4: Package
) -> None:
    await run_alerts(session, NOW)
    session.expire_all()
    package = (await session.execute(select(Package).where(Package.number == 4))).scalar_one()
    assert package.health == "red"


async def test_a_closed_package_raises_nothing(
    session: AsyncSession, plan_for_package_4: Package
) -> None:
    plan_for_package_4.status = "settled"
    await session.commit()
    await run_alerts(session, NOW)
    assert await plan_alerts(session) == []


# --- dashboard and export ------------------------------------------------------------------------


async def test_steps_due_within_30_days_are_milestones(
    session: AsyncSession, plan_for_package_4: Package, make_client_for, monkeypatch
) -> None:
    monkeypatch.setattr("app.routers.dashboard.today_local", lambda: date(2026, 11, 15))
    c = await make_client_for(session, "viewer")
    data = (await c.get("/api/v1/dashboard/summary")).json()
    steps = [m for m in data["milestones"] if m["kind"] == "plan_step"]
    assert [(m["date"], m["days_left"]) for m in steps] == [("2026-11-18", 3)]  # step 1 is past
    assert steps[0]["title"].startswith("Bước 2 kế hoạch") and steps[0]["package_number"] == 4

    step = (await session.execute(select(PlanStep).where(PlanStep.step_no == 2))).scalar_one()
    step.status = "done"
    await session.commit()
    again = (await c.get("/api/v1/dashboard/summary")).json()
    assert [m for m in again["milestones"] if m["kind"] == "plan_step"] == []


async def test_alert_export_names_the_new_type(
    session: AsyncSession, plan_for_package_4: Package, make_client_for
) -> None:
    import io

    from openpyxl import load_workbook

    await run_alerts(session, NOW)
    c = await make_client_for(session, "viewer")
    ws = load_workbook(io.BytesIO((await c.get("/api/v1/export/alerts.xlsx")).content)).active
    types = {row[1].value for row in ws.iter_rows(min_row=5)}
    assert "Bước kế hoạch quá hạn" in types
