"""M8: Excel import of contract goods lines (SPEC 4.13) - dry run, validation, all-or-nothing."""

import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import AuditLog, Contract, ContractItem, Package
from app.seed.project import seed_project
from app.services.import_items import (
    HEADERS,
    ImportFileError,
    parse_number,
    parse_workbook,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
URL = "/api/v1/import/contract-items"
WRITERS = {"admin", "director", "procurement", "cost"}


def workbook(rows: list[list], headers: list[str] | None = None, *, lead: int = 0) -> bytes:
    wb = Workbook()
    ws = wb.active
    for _ in range(lead):
        ws.append(["Danh mục hàng hóa hợp đồng"])
    ws.append(headers if headers is not None else list(HEADERS))
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


GOOD = [
    [1, "Máy tính xách tay", "Cái", 10, 15_000_000, 150_000_000, 36, "x", "Việt Nam", None],
    [2, "Máy in", "Cái", 5, 4_000_000, None, 12, None, None, "Màu"],
]


# --- parser ------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (5, Decimal(5)),
        (2.5, Decimal("2.5")),
        ("1.234.567", Decimal(1234567)),
        ("1.234,5", Decimal("1234.5")),
        ("12,5", Decimal("12.5")),
        ("3.5", Decimal("3.5")),
        ("15.000.000 đ", Decimal(15000000)),
        (" ", None),
        (None, None),
    ],
)
def test_parse_number(raw: object, expected: Decimal | None) -> None:
    assert parse_number(raw) == expected


@pytest.mark.parametrize("raw", ["abc", "1,2,3", True])
def test_parse_number_rejects_text(raw: object) -> None:
    with pytest.raises(ValueError):
        parse_number(raw)


def test_parses_rows_and_computes_a_missing_amount() -> None:
    result = parse_workbook(workbook(GOOD))
    assert result.errors == [] and result.rows_seen == 2
    first, second = result.items
    assert first.name == "Máy tính xách tay" and first.requires_calibration is True
    assert second.amount == Decimal(20_000_000) and second.requires_calibration is False
    assert second.notes == "Màu" and second.warranty_months == 12
    assert result.total_amount == Decimal(170_000_000)


def test_finds_the_header_below_a_title_and_accepts_unaccented_headers() -> None:
    headers = ["stt", "ten hang hoa", "dvt", "so luong", "don gia"]
    result = parse_workbook(workbook([[1, "Bàn", "Cái", "2", "1.000.000"]], headers, lead=2))
    assert result.errors == [] and result.items[0].amount == Decimal(2_000_000)
    assert result.items[0].row == 4  # sheet row, 1-based: 2 title rows + header + data


def test_row_errors_carry_the_sheet_row_and_field() -> None:
    rows = [
        [1, "Tốt", "Cái", 1, 100, 100, 12, None, None, None],
        [2, None, "Cái", 1, 100, 100, None, None, None, None],  # no name
        [3, "Sai số", "Cái", "hai", 100, None, None, None, None, None],  # not a number
        [4, "Âm", "Cái", -1, 100, None, None, None, None, None],
        [5, "Lệch", "Cái", 2, 100, 500, None, None, None, None],  # amount != 2 x 100
        [6, "Bảo hành lẻ", "Cái", 1, 100, 100, 1.5, None, None, None],
    ]
    result = parse_workbook(workbook(rows))
    assert len(result.items) == 1
    by_row = {(e.row, e.field) for e in result.errors}
    assert by_row == {
        (3, "name"),
        (4, "quantity"),
        (5, "quantity"),
        (6, "amount"),
        (7, "warranty_months"),
    }


def test_blank_rows_are_skipped_and_one_dong_of_rounding_is_tolerated() -> None:
    rows = [[None] * 10, [1, "Cáp", "m", 3, "333,33", 1000, None, None, None, None]]
    result = parse_workbook(workbook(rows))
    assert result.errors == [] and result.rows_seen == 1


@pytest.mark.parametrize(
    "content",
    [b"not a workbook", b"", workbook([], headers=["a", "b", "c"]), workbook([], headers=[])],
)
def test_unreadable_or_headerless_files_are_rejected(content: bytes) -> None:
    with pytest.raises(ImportFileError):
        parse_workbook(content)


def test_oversized_files_are_rejected() -> None:
    with pytest.raises(ImportFileError, match="quá lớn"):
        parse_workbook(b"0" * (5 * 1024 * 1024 + 1))


# --- API ---------------------------------------------------------------------------------------


@pytest.fixture
async def contract(session: AsyncSession) -> Contract:
    await seed_project(session)
    package = (await session.execute(select(Package).where(Package.number == 4))).scalar_one()
    found = (
        await session.execute(select(Contract).where(Contract.package_id == package.id))
    ).scalar_one()
    found.value = Decimal(170_000_000)
    await session.commit()
    return found


def upload(content: bytes, name: str = "hang.xlsx") -> dict:
    return {"file": (name, content, XLSX)}


async def items_of(session: AsyncSession, contract: Contract) -> list[ContractItem]:
    stmt = (
        select(ContractItem)
        .where(ContractItem.contract_id == contract.id)
        .order_by(ContractItem.line_no)
        .execution_options(populate_existing=True)
    )
    return list((await session.execute(stmt)).scalars().all())


async def test_dry_run_reports_everything_and_writes_nothing(
    session: AsyncSession, contract: Contract, make_client_for
) -> None:
    c = await make_client_for(session, "procurement")
    r = await c.post(URL, params={"contract_id": str(contract.id)}, files=upload(workbook(GOOD)))
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["dry_run"] is True and data["imported"] == 0
    assert (data["rows_total"], data["rows_valid"], data["errors"]) == (2, 2, [])
    assert data["total_amount"] == 170_000_000 and data["contract_value"] == 170_000_000
    assert data["matches_contract_value"] is True
    assert [p["line_no"] for p in data["preview"]] == [1, 2]
    assert await items_of(session, contract) == []


async def test_commit_writes_the_lines_and_audits(
    session: AsyncSession, contract: Contract, make_client_for
) -> None:
    c = await make_client_for(session, "procurement")
    r = await c.post(
        URL,
        params={"contract_id": str(contract.id), "commit": "true"},
        files=upload(workbook(GOOD)),
    )
    assert r.status_code == 200 and r.json()["imported"] == 2
    items = await items_of(session, contract)
    assert [(i.line_no, i.name, i.amount) for i in items] == [
        (1, "Máy tính xách tay", Decimal(150_000_000)),
        (2, "Máy in", Decimal(20_000_000)),
    ]
    assert items[0].requires_calibration is True and items[0].origin == "Việt Nam"
    log = (
        await session.execute(
            select(AuditLog).where(AuditLog.entity_type == "contract_items_import")
        )
    ).scalar_one()
    assert log.changes["rows"] == 2 and log.changes["replace"] is False


async def test_append_continues_the_numbering_and_replace_starts_over(
    session: AsyncSession, contract: Contract, make_client_for
) -> None:
    c = await make_client_for(session, "procurement")
    params = {"contract_id": str(contract.id), "commit": "true"}
    await c.post(URL, params=params, files=upload(workbook(GOOD)))
    await c.post(URL, params=params, files=upload(workbook(GOOD)))
    assert [i.line_no for i in await items_of(session, contract)] == [1, 2, 3, 4]

    r = await c.post(URL, params={**params, "replace": "true"}, files=upload(workbook(GOOD[:1])))
    assert r.json()["imported"] == 1
    assert [i.name for i in await items_of(session, contract)] == ["Máy tính xách tay"]


async def test_a_mismatch_with_the_contract_value_is_reported_not_blocked(
    session: AsyncSession, contract: Contract, make_client_for
) -> None:
    c = await make_client_for(session, "procurement")
    r = await c.post(
        URL, params={"contract_id": str(contract.id)}, files=upload(workbook(GOOD[:1]))
    )
    data = r.json()
    assert data["total_amount"] == 150_000_000 and data["matches_contract_value"] is False

    contract.value = None
    await session.commit()
    again = await c.post(
        URL, params={"contract_id": str(contract.id)}, files=upload(workbook(GOOD[:1]))
    )
    assert again.json()["matches_contract_value"] is None


async def test_commit_with_errors_is_refused_and_nothing_is_written(
    session: AsyncSession, contract: Contract, make_client_for
) -> None:
    c = await make_client_for(session, "procurement")
    rows = [*GOOD, [3, None, "Cái", 1, 1, None, None, None, None, None]]
    dry = await c.post(URL, params={"contract_id": str(contract.id)}, files=upload(workbook(rows)))
    assert dry.status_code == 200 and dry.json()["errors"][0]["row"] == 4
    assert dry.json()["rows_valid"] == 2

    r = await c.post(
        URL,
        params={"contract_id": str(contract.id), "commit": "true"},
        files=upload(workbook(rows)),
    )
    assert (r.status_code, r.json()["error"]["code"]) == (422, "import_has_errors")
    assert r.json()["error"]["details"][0]["field"] == "name"
    assert await items_of(session, contract) == []  # the two good rows were not written either


async def test_bad_requests(session: AsyncSession, contract: Contract, make_client_for) -> None:
    c = await make_client_for(session, "procurement")
    params = {"contract_id": str(contract.id)}
    r = await c.post(URL, params=params, files=upload(b"zip", "hang.csv"))
    assert r.status_code == 422  # wrong extension
    r = await c.post(URL, params=params, files=upload(b"junk", "hang.xlsx"))
    assert (r.status_code, r.json()["error"]["code"]) == (422, "invalid_import_file")
    ghost = "00000000-0000-0000-0000-000000000000"
    r = await c.post(URL, params={"contract_id": ghost}, files=upload(workbook(GOOD)))
    assert r.status_code == 404
    assert (await c.post(URL, params=params)).status_code == 422  # no file
    assert (await c.post(URL, files=upload(workbook(GOOD)))).status_code == 422  # no contract
    empty = workbook([])
    r = await c.post(URL, params={**params, "commit": "true"}, files=upload(empty))
    assert r.status_code == 422 and "không có dòng" in r.json()["error"]["message"]


async def test_template_has_the_expected_headers(session: AsyncSession, make_client_for) -> None:
    c = await make_client_for(session, "cost")
    r = await c.get(f"{URL}/template.xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [cell.value for cell in ws[4]] == list(HEADERS)
    # the template itself imports cleanly
    assert parse_workbook(r.content).errors == []


async def test_role_matrix(
    session: AsyncSession, contract: Contract, make_client_for, client
) -> None:
    from tests.conftest import ALL_ROLES

    params = {"contract_id": str(contract.id)}
    assert (await client.post(URL, params=params, files=upload(workbook(GOOD)))).status_code == 401
    for role in ALL_ROLES:
        c = await make_client_for(session, role)
        r = await c.post(URL, params=params, files=upload(workbook(GOOD)))
        assert r.status_code == (200 if role in WRITERS else 403), role
        t = await c.get(f"{URL}/template.xlsx")
        assert t.status_code == (200 if role in WRITERS else 403), role
